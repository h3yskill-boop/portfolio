"""
Excel Exporter Module.
Handles exporting scraped product data into a formatted Excel (.xlsx) file using pandas and openpyxl.
"""

from datetime import datetime
import logging
from typing import List
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import OUTPUT_FILE_PREFIX
from parser import Product


class ExcelReportBuilder:
    """
    Class responsible for generating styled Excel reports from product data.
    """

    def __init__(self, products: List[Product]):
        self.products = products

    def build_dataframe(self) -> pd.DataFrame:
        """
        Converts a list of Product models into a structured pandas DataFrame.
        """
        data = [product.model_dump() for product in self.products]
        df = pd.DataFrame(data)

        # Rename columns to human-readable names for the Excel report
        df.rename(
            columns={
                "title": "Product Title",
                "price": "Price ($)",
                "old_price": "Old Price ($)",
                "discount_percent": "Discount (%)",
                "availability": "Availability",
                "rating": "Rating (Stars)",
                "reviews_count": "Reviews Count",
                "image_url": "Image Link",
                "product_url": "Product Link",
            },
            inplace=True,
        )

        # Reorder columns logically
        columns_order = [
            "Product Title",
            "Price ($)",
            "Old Price ($)",
            "Discount (%)",
            "Availability",
            "Rating (Stars)",
            "Reviews Count",
            "Product Link",
            "Image Link",
        ]
        return df[columns_order]

    def export(self) -> str:
        """
        Generates and styles the Excel spreadsheet, then saves it to disk.
        """
        df = self.build_dataframe()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_name = f"{OUTPUT_FILE_PREFIX}_{timestamp}.xlsx"

        # Create Excel writer using openpyxl engine
        with pd.ExcelWriter(file_name, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Competitor Analysis")
            workbook = writer.book
            worksheet = writer.sheets["Competitor Analysis"]

            # Enable AutoFilter across all columns
            worksheet.auto_filter.ref = worksheet.dimensions

            # Styles Definition
            header_fill = PatternFill(
                start_color="1F4E78", end_color="1F4E78", fill_type="solid"
            )  # Dark Blue
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            data_font = Font(name="Calibri", size=11)
            thin_border = Border(
                left=Side(style="thin", color="D9D9D9"),
                right=Side(style="thin", color="D9D9D9"),
                top=Side(style="thin", color="D9D9D9"),
                bottom=Side(style="thin", color="D9D9D9"),
            )

            # Apply styling to Header Row
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )

            # Apply styling and formatting to Data Rows
            for row in worksheet.iter_rows(
                min_row=2, max_row=len(df) + 1, min_col=1, max_col=len(df.columns)
            ):
                for cell in row:
                    cell.font = data_font
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center")

            # Apply Number Formatting for prices and percentages
            # Price columns: B (2), C (3)
            for col_idx in [2, 3]:
                for cell in worksheet.iter_cols(
                    min_col=col_idx,
                    max_col=col_idx,
                    min_row=2,
                    max_row=len(df) + 1,
                ):
                    for c in cell:
                        c.number_format = "$#,##0.00"

            # Discount column: D (4)
            for cell in worksheet.iter_cols(
                min_col=4, max_col=4, min_row=2, max_row=len(df) + 1
            ):
                for c in cell:
                    c.number_format = '0.0"%"'

            # Auto-adjust column widths based on max string length
            for col in worksheet.columns:
                max_length = 0
                column_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                # Set dynamic width with safety padding (capped at 50 for URLs)
                adjusted_width = min(max(max_length + 3, 12), 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width

        logging.info(f"Report successfully saved to: {file_name}")
        return file_name